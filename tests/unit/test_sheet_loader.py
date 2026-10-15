"""Testes unitários para o leitor de planilhas (sheet_loader)."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from core.models import SheetFormatError
from storage.sheet_loader import load_jobs, parse_fit_score, split_list


def test_vagas_sample_csv() -> None:
    csv_path = Path("tests/fixtures/vagas_sample.csv")
    jobs = load_jobs(csv_path, only_todo=True)

    # 5 linhas "A Fazer" no fixture CSV
    # 1. Audax Soluções
    # 2. CI&T
    # 3. Sinqia / Evertec
    # 4. WithProtégé AI
    # 5. Empresa Quebrada
    assert len(jobs) == 5

    j_audax = next(j for j in jobs if "audax" in j.company.lower())
    assert j_audax.fit_score == 95
    assert j_audax.language == "pt"
    assert "Cognition" in j_audax.base_used
    assert "Base Português" in j_audax.base_used
    assert "Next.js 16" in j_audax.stack
    assert "Supabase" in j_audax.stack
    assert j_audax.url == "https://jobs.lever.co/audax/123"

    j_cit = next(j for j in jobs if "ci&t" in j.company.lower())
    assert j_cit.fit_score == 84
    assert j_cit.language == "en"
    assert "Base Inglês" in j_cit.base_used
    assert j_cit.url == "https://boards.greenhouse.io/cit/jobs/456"

    j_protege = next(j for j in jobs if "withprotege" in j.id)
    assert j_protege.language == "en"
    # "Ver Vaga" no CSV perde a URL na exportação, vira url=None sem derrubar
    assert j_protege.url is None

    j_quebrada = next(j for j in jobs if "quebrada" in j.company.lower())
    assert j_quebrada.fit_score is None
    assert j_quebrada.resume_label is None
    assert j_quebrada.url is None

    # Todos os IDs devem ser únicos e não vazios
    ids = [j.id for j in jobs]
    assert len(ids) == len(set(ids))
    assert all(len(i) > 0 for i in ids)


def test_xlsx_hyperlink_and_header(tmp_path: Path) -> None:
    """Verifica extração de hiperlinks no formato .xlsx e detecção de cabeçalho com linhas vazias."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None

    # Linha 1: Título
    ws.cell(row=1, column=2, value="Planilha de Vagas")
    # Linha 2 e 3 em branco

    # Linha 4: Cabeçalho (coluna A vazia, cabeçalho na coluna B..K)
    headers = [
        "Data", "Empresa", "Cargo", "Modalidade", "Link da Vaga",
        "Currículo Otimizado", "Fit Score", "Base Utilizada", "Stack Principal", "Status da Candidatura"
    ]
    for col_idx, h in enumerate(headers, start=2):
        ws.cell(row=4, column=col_idx, value=h)

    # Linha 5: Linha de dados com hiperlink na célula
    ws.cell(row=5, column=2, value="01/10/2026")
    ws.cell(row=5, column=3, value="Empresa Hyperlink")
    ws.cell(row=5, column=4, value="Staff Engineer")
    ws.cell(row=5, column=5, value="Remoto")
    cell_link = ws.cell(row=5, column=6, value="Ver Vaga")
    cell_link.hyperlink = "https://jobs.lever.co/hyperlink/1"
    ws.cell(row=5, column=7, value="Currículo Otimizado - Hyperlink")
    ws.cell(row=5, column=8, value="92%")
    ws.cell(row=5, column=9, value="Cognition")
    ws.cell(row=5, column=10, value="Python, Go")
    ws.cell(row=5, column=11, value="A Fazer")

    # Linha 6: Linha com status diferente ("Candidatado")
    ws.cell(row=6, column=3, value="Empresa Ignorada")
    ws.cell(row=6, column=4, value="Dev")
    ws.cell(row=6, column=11, value="Candidatado")

    xlsx_path = tmp_path / "teste_vagas.xlsx"
    wb.save(xlsx_path)

    jobs = load_jobs(xlsx_path, only_todo=True)
    assert len(jobs) == 1
    assert jobs[0].company == "Empresa Hyperlink"
    assert jobs[0].url == "https://jobs.lever.co/hyperlink/1"
    assert jobs[0].fit_score == 92


def test_missing_header_raises_sheet_format_error(tmp_path: Path) -> None:
    """Planilha sem Empresa e Cargo deve levantar SheetFormatError."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(row=1, column=1, value="Invalida")
    p = tmp_path / "invalid.xlsx"
    wb.save(p)

    with pytest.raises(SheetFormatError):
        load_jobs(p)


def test_fit_score_parsing_variations() -> None:
    assert parse_fit_score("95%") == 95
    assert parse_fit_score("95") == 95
    assert parse_fit_score("0.95") == 95
    assert parse_fit_score("#ERROR!") is None
    assert parse_fit_score("") is None
    assert parse_fit_score("invalid") is None
    assert parse_fit_score("150%") is None


def test_split_list() -> None:
    assert split_list("A / B / C", " / ") == ["A", "B", "C"]
    assert split_list("A, B, C", ",") == ["A", "B", "C"]
    assert split_list("#ERROR!", ",") == []
    assert split_list("", ",") == []


def test_google_sheet_detection_and_extraction() -> None:
    from storage.sheet_loader import extract_spreadsheet_id, is_google_sheet_source

    url_edit = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#gid=0"
    url_sharing = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit?usp=sharing"
    raw_id = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"

    assert is_google_sheet_source(url_edit) is True
    assert is_google_sheet_source(url_sharing) is True
    assert is_google_sheet_source(raw_id) is True
    assert is_google_sheet_source("data/vagas.xlsx") is False

    assert extract_spreadsheet_id(url_edit) == raw_id
    assert extract_spreadsheet_id(url_sharing) == raw_id
    assert extract_spreadsheet_id(raw_id) == raw_id


def test_download_google_sheet_http_export(tmp_path: Path) -> None:
    import respx

    from storage.sheet_loader import download_google_sheet

    fake_id = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
    export_url = f"https://docs.google.com/spreadsheets/d/{fake_id}/export?format=xlsx"
    target_file = tmp_path / "downloaded.xlsx"

    fake_xlsx_bytes = b"PK\x03\x04" + b"\x00" * 100

    with respx.mock:
        respx.get(export_url).respond(
            status_code=200,
            content=fake_xlsx_bytes,
            headers={"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        )
        saved = download_google_sheet(fake_id, dest_path=target_file)
        assert saved == target_file
        assert target_file.exists()
        assert target_file.read_bytes() == fake_xlsx_bytes


def test_load_jobs_include_blank_status(tmp_path: Path) -> None:
    """Verifica comportamento com status em branco quando include_blank é ativado ou desativado."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None

    headers = [
        "Data", "Empresa", "Cargo", "Modalidade", "Link da Vaga",
        "Currículo Otimizado", "Fit Score", "Base Utilizada", "Stack Principal", "Status da Candidatura"
    ]
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)

    # Linha 2: "A Fazer"
    ws.cell(row=2, column=2, value="Empresa Todo")
    ws.cell(row=2, column=3, value="Engenheiro")
    ws.cell(row=2, column=5, value="https://jobs.lever.co/empresa/1")
    ws.cell(row=2, column=10, value="A Fazer")

    # Linha 3: Status vazio ("")
    ws.cell(row=3, column=2, value="Empresa Blank")
    ws.cell(row=3, column=3, value="Tech Lead")
    ws.cell(row=3, column=5, value="https://jobs.lever.co/empresa/2")
    ws.cell(row=3, column=10, value="")

    p = tmp_path / "vagas_blank.xlsx"
    wb.save(p)

    # Com include_blank=False (padrão): só carrega a linha 2
    jobs_strict = load_jobs(p, only_todo=True, include_blank=False)
    assert len(jobs_strict) == 1
    assert jobs_strict[0].company == "Empresa Todo"

    # Com include_blank=True: carrega ambas
    jobs_all = load_jobs(p, only_todo=True, include_blank=True)
    assert len(jobs_all) == 2

