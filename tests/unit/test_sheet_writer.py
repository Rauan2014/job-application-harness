"""Testes unitários para o escritor de planilhas (sheet_writer)."""

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.models import ApplicationResult, TaskStatus
from storage.sheet_writer import write_results


def _get_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_source_untouched_and_status_mapped(tmp_path: Path) -> None:
    source_csv = Path("tests/fixtures/vagas_sample.csv")
    initial_hash = _get_hash(source_csv)

    results = [
        ApplicationResult(
            job_id="audax-solucoes-ai-developer-saas",
            status=TaskStatus.COMPLETED,
            dry_run=True,
            submitted=False,
            reason="Simulação dry-run bem-sucedida",
        ),
        ApplicationResult(
            job_id="ci-t-job-31910-ai-engineer-brazil",
            status=TaskStatus.COMPLETED,
            dry_run=False,
            submitted=True,
            reason="Envio confirmado com sucesso",
        ),
        ApplicationResult(
            job_id="sinqia-evertec-senior-python-developer",
            status=TaskStatus.NEEDS_REVIEW,
            reason="CAPTCHA detectado na página",
        ),
        ApplicationResult(
            job_id="withprotege-ai-ai-agent-engineer",
            status=TaskStatus.SKIPPED,
            reason="Link ausente na planilha",
        ),
    ]

    from storage.sheet_loader import load_jobs

    jobs = {j.id: j for j in load_jobs(source_csv, only_todo=False, include_blank=True)}
    out_xlsx = tmp_path / "planilha_resultado.xlsx"
    write_results(source_csv, results, out_xlsx, jobs=jobs)

    # 1. Arquivo de origem não foi tocado
    assert _get_hash(source_csv) == initial_hash

    # 2. Arquivo de saída existe e contém os dados atualizados
    assert out_xlsx.exists()
    wb = load_workbook(out_xlsx)
    ws = wb.active
    assert ws is not None

    # Mapeamento de cabeçalhos
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    assert "Motivo" in headers
    assert "Processado em" in headers

    status_col = headers.index("Status da Candidatura") + 1
    reason_col = headers.index("Motivo") + 1
    proc_col = headers.index("Processado em") + 1

    # Linha 5 (Audax) -> Pronta (dry-run)
    assert ws.cell(row=5, column=status_col).value == "Pronta (dry-run)"
    assert "Simulação dry-run" in str(ws.cell(row=5, column=reason_col).value)
    assert ws.cell(row=5, column=proc_col).value is not None

    # Linha 6 (CI&T) -> Candidatado
    assert ws.cell(row=6, column=status_col).value == "Candidatado"

    # Linha 7 (Sinqia) -> Revisar
    assert ws.cell(row=7, column=status_col).value == "Revisar"
    assert "CAPTCHA" in str(ws.cell(row=7, column=reason_col).value)

    # Linha 8 (WithProtégé) -> Ignorada
    assert ws.cell(row=8, column=status_col).value == "Ignorada"

    # Linha 9 (Empresa Concluída - não estava na fila de resultados) -> Preservada
    assert ws.cell(row=9, column=status_col).value == "Concluído"


def test_cannot_overwrite_source(tmp_path: Path) -> None:
    source = tmp_path / "mesmo.xlsx"
    source.write_text("dummy", encoding="utf-8")
    with pytest.raises(ValueError, match="out_path não pode ser a planilha original"):
        write_results(source, [], source, jobs={})


def test_write_results_with_blank_status_jobs(tmp_path: Path) -> None:
    from core.models import Job

    source_csv = Path("tests/fixtures/vagas_sample.csv")
    job_blank = Job(
        id="vaga-em-branco",
        row_index=5,
        company="Audax Soluções",
        title="AI Developer",
        sheet_status="",
    )
    jobs_map = {job_blank.id: job_blank}
    result = ApplicationResult(
        job_id="vaga-em-branco",
        status=TaskStatus.COMPLETED,
        dry_run=True,
        submitted=False,
        reason="Vaga com status em branco processada com sucesso",
    )
    out_xlsx = tmp_path / "out_blank.xlsx"
    write_results(source_csv, [result], out_xlsx, jobs=jobs_map)

    wb = load_workbook(out_xlsx)
    ws = wb.active
    assert ws is not None

    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    status_col = headers.index("Status da Candidatura") + 1
    reason_col = headers.index("Motivo") + 1
    assert ws.cell(row=5, column=status_col).value == "Pronta (dry-run)"
    assert ws.cell(row=5, column=reason_col).value == "Vaga com status em branco processada com sucesso"


def test_write_results_updates_resolved_link_url(tmp_path: Path) -> None:
    """Garante que write_results também corrige a coluna 'Link da Vaga' quando job.url foi resolvida."""
    from core.models import Job

    source_csv = Path("tests/fixtures/vagas_sample.csv")
    resolved_url = "https://jobs.lever.co/ciandt/56f9af4b-0b7c-4f99-be77-e51a2859abdd"
    job = Job(
        id="withprotege-ai-ai-agent-engineer",
        row_index=8,
        company="WithProtégé AI",
        title="AI Agent Engineer",
        url=resolved_url,
    )
    res = ApplicationResult(
        job_id=job.id,
        status=TaskStatus.COMPLETED,
        dry_run=False,
        submitted=True,
        reason="Candidatura enviada e confirmada (lever).",
    )
    out_xlsx = tmp_path / "out_resolved_url.xlsx"
    write_results(source_csv, [res], out_xlsx, jobs={job.id: job})

    wb = load_workbook(out_xlsx)
    ws = wb.active
    assert ws is not None

    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    link_col = headers.index("Link da Vaga") + 1
    status_col = headers.index("Status da Candidatura") + 1
    assert ws.cell(row=8, column=status_col).value == "Candidatado"
    assert ws.cell(row=8, column=link_col).value == resolved_url
    assert ws.cell(row=8, column=link_col).hyperlink is not None
    assert ws.cell(row=8, column=link_col).hyperlink.target == resolved_url

