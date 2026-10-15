"""Sincroniza automaticamente os currículos otimizados linkados na planilha do Google Sheets."""

from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from config.settings import load_settings
from storage.drive_client import build_service

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("sync_resumes")


def sync_sheet_resumes(sheet_path: Path, cache_dir: Path) -> int:
    cache_dir.mkdir(parents=True, exist_ok=True)
    settings = load_settings()
    svc = build_service(settings.GOOGLE_SERVICE_ACCOUNT_FILE, settings.GOOGLE_OAUTH_CLIENT_FILE)

    wb = openpyxl.load_workbook(str(sheet_path), data_only=False)
    ws = wb.active
    downloaded = 0

    for r in range(1, ws.max_row + 1):
        cell_val = str(ws.cell(row=r, column=7).value or "")
        m = re.search(r'HYPERLINK\("([^"]+)",\s*"([^"]+)"\)', cell_val)
        if not m:
            continue
        url, label = m.group(1), m.group(2)
        doc_id_m = re.search(r'/d/([a-zA-Z0-9_-]+)', url)
        if not doc_id_m:
            continue

        doc_id = doc_id_m.group(1)
        safe_label = "".join(c for c in label if c not in '<>:"/\\|?*').strip()
        pdf_path = cache_dir / f"{safe_label}.pdf"

        # Baixa se não existir ou se for o arquivo mock pequeno anterior (< 5 KB)
        if not pdf_path.exists() or pdf_path.stat().st_size < 5000:
            logger.info("Baixando currículo otimizado '%s' (doc_id=%s)...", safe_label, doc_id)
            try:
                pdf_bytes = svc.files().export(fileId=doc_id, mimeType="application/pdf").execute()
                pdf_path.write_bytes(pdf_bytes)
                logger.info("Salvo: %s (%d bytes)", pdf_path.name, len(pdf_bytes))
                downloaded += 1
            except Exception as exc:
                logger.error("Erro ao baixar doc %s: %s", doc_id, exc)
        else:
            logger.info("Já presente em cache: %s (%d bytes)", pdf_path.name, pdf_path.stat().st_size)

    wb.close()
    return downloaded


if __name__ == "__main__":
    count = sync_sheet_resumes(Path("cache/vagas_sheets_latest.xlsx"), Path("cache/resumes"))
    print(f"\nDownload concluído! Total de novos currículos sincronizados: {count}")
